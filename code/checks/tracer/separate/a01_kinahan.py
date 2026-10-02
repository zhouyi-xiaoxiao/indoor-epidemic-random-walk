import sys, json, numpy as np
import openpyxl
wb=openpyxl.load_workbook('../../../bsc_validation2/a3_tracer_physics/data/raw/kinahan_777_inflight.xlsx',read_only=True,data_only=True)
print(wb.sheetnames)
ws=wb['FWD'] if 'FWD' in wb.sheetnames else wb[[s for s in wb.sheetnames if 'FWD' in s][0]]
rows=[]
for i,r in enumerate(ws.iter_rows(min_row=1,max_row=60,max_col=48,values_only=True)):
    rows.append(r)
hdr=[i for i,r in enumerate(rows) if r[0] and str(r[0]).startswith('Release')][0]
h=rows[hdr]; print(h[:8]); sens=[str(c).split()[-1] if c else None for c in h[7:47]]
print(sens)
out={}
hdr2=[i for i,r in enumerate(rows) if r[0] and str(r[0]).startswith("Release")][1]
for r in rows[hdr+2:hdr2]:
    if r[0] is None or r[2] not in('B','C'): 
        continue
    print(r[0],r[1],r[2],r[3],r[4])
    if r[2]=='B' and str(r[4]).strip()=='No':
        out.setdefault(str(r[0]).strip(),[]).append([np.nan if v is None or v=='' else float(v) for v in r[7:47]])
res={}
for k,v in out.items():
    a=np.array(v); print(k,a.shape)
    import warnings; warnings.simplefilter('ignore')
    res[k]=dict(zip(sens,np.nanmean(a,axis=0).tolist()))
json.dump(res,open('kinahan_fwd_means.json','w'),indent=0)
for k in res: print(k,{s:(None if v!=v else round(v,1)) for s,v in res[k].items()})
