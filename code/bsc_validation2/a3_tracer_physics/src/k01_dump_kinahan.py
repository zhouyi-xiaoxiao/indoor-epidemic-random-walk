"""Dump the small summary sheets of the Kinahan et al. figshare workbooks to CSV (values only)."""
import sys,os,csv
sys.path.insert(0,os.path.join(os.path.dirname(__file__),'..','.pylib'))
import openpyxl
root=os.path.join(os.path.dirname(__file__),'..','data')
for af in ('777','767'):
    wb=openpyxl.load_workbook(os.path.join(root,'raw',f'kinahan_{af}_inflight.xlsx'),read_only=True,data_only=True)
    for ws in wb.worksheets:
        ws.reset_dimensions()
        print(af,ws.title)
        if 'raw' in ws.title.lower() or 'data' in ws.title.lower(): 
            # only peek
            for i,row in enumerate(ws.iter_rows(values_only=True)):
                if i<4: print('   ',[c for c in row[:14]])
                else: break
            continue
        out=os.path.join(root,'derived',f'kinahan_{af}_{ws.title.replace(" ","_").replace("/","-")}.csv')
        os.makedirs(os.path.dirname(out),exist_ok=True)
        with open(out,'w',newline='') as f:
            w=csv.writer(f)
            for row in ws.iter_rows(values_only=True):
                w.writerow(['' if c is None else c for c in row])
